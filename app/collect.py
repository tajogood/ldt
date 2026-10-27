"""원천별 수집기 + 표준 테이블로 정규화.

흐름: fetch(원천) → 원본 JSON을 data/raw/<id>/<날짜>.json 에 보관 → map 규칙으로 정규화 → upsert
원천 응답 필드명은 신청 후 실제 호출로 확인해야 하므로, sources.yaml 의 후보 목록으로 맞춘다.
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterator

import httpx
import yaml

from . import db, geo
from .settings import CONFIG_DIR, DATA_GO_KR_KEY, KEEP_RAW_FILES, MAFRA_KEY, RAW_DIR, ROOT
from .species import disease_code, species_code

MAFRA_BASE = "http://211.237.50.150:7080/openapi"
TIMEOUT = httpx.Timeout(30.0)


# ── 설정 ────────────────────────────────────────────────────────────────
def load_sources() -> list[dict]:
    with open(CONFIG_DIR / "sources.yaml", encoding="utf-8") as f:
        return yaml.safe_load(f)["sources"]


def get_source(sid: str) -> dict:
    for s in load_sources():
        if s["id"] == sid:
            return s
    raise KeyError(f"sources.yaml 에 '{sid}' 가 없습니다")


def is_configured(src: dict) -> bool:
    if src.get("enabled") is False:
        return False
    p = src["provider"]
    if p == "mafra":
        return bool(src.get("service_id")) and bool(MAFRA_KEY)
    if p in ("datagokr", "odcloud", "kma_asos"):
        return bool(src.get("endpoint")) and bool(DATA_GO_KR_KEY)
    if p == "file":
        return (ROOT / src["path"]).exists()
    if p == "kahis_visit":
        return True
    if p == "birds":
        return any(x.suffix.lower() in (".csv", ".xlsx") and ("철새" in x.name or "migrat" in x.name.lower())
                   for x in (ROOT / "data" / "files").glob("*"))
    return False


# ── 공통 유틸 ──────────────────────────────────────────────────────────
def find_records(obj: Any) -> list[dict]:
    """응답 JSON 어디에 있든 가장 큰 dict 리스트를 레코드로 본다."""
    best: list[dict] = []

    def walk(o: Any) -> None:
        nonlocal best
        if isinstance(o, list):
            if o and all(isinstance(x, dict) for x in o) and len(o) > len(best):
                best = o
            for x in o:
                walk(x)
        elif isinstance(o, dict):
            for v in o.values():
                walk(v)

    walk(obj)
    if not best and isinstance(obj, dict):
        # item 이 1건일 때 dict 로 오는 data.go.kr 응답 대응
        items = obj.get("response", {}).get("body", {}).get("items", {})
        if isinstance(items, dict) and isinstance(items.get("item"), dict):
            best = [items["item"]]
    return best


def total_count(obj: Any) -> int | None:
    m = re.search(r'"(totalCount|totalCnt|matchCount|list_total_count)"\s*:\s*"?(\d+)', json.dumps(obj))
    return int(m.group(2)) if m else None


def pick(rec: dict, candidates: list[str] | None) -> Any:
    if not candidates:
        return None
    lower = {str(k).lower(): v for k, v in rec.items() if k is not None}
    for c in candidates:
        v = rec.get(c, lower.get(c.lower()))
        if v not in (None, "", " "):
            return v.strip() if isinstance(v, str) else v
    return None


def norm_date(v: Any) -> str | None:
    if v in (None, ""):
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    s = re.sub(r"[^0-9]", "", str(v))
    if len(s) >= 8:
        return f"{s[:4]}-{s[4:6]}-{s[6:8]}"
    if len(s) == 6:
        return f"{s[:4]}-{s[4:6]}-01"
    if len(s) == 4:
        return f"{s}-01-01"
    return None


def num(v: Any) -> float | None:
    if v in (None, ""):
        return None
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return None


def make_uid(source: str, parts: list[Any]) -> str:
    h = hashlib.sha1("|".join("" if p is None else str(p) for p in parts).encode()).hexdigest()[:16]
    return f"{source}:{h}"


def save_raw(sid: str, payload: Any) -> Path | None:
    """페이지 단위 원본 JSON 파일(선택). DB raw_record 에는 run() 이 레코드 단위로 항상 저장한다."""
    if not KEEP_RAW_FILES:
        return None
    d = RAW_DIR / sid
    d.mkdir(parents=True, exist_ok=True)
    p = d / f"{datetime.now():%Y%m%d_%H%M%S}.json"
    p.write_text(json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8")
    return p


# ── fetchers ────────────────────────────────────────────────────────────
def iter_values(it: dict) -> list[str]:
    """iterate 설정 → 파라미터 값 목록. kind: month(YYYYMM) | year(YYYY) | list."""
    kind = it.get("kind", "list")
    today = date.today()
    if kind == "list":
        return [str(v) for v in it["values"]]
    if kind == "year":
        y0 = int(str(it["from"])[:4])
        y1 = int(str(it.get("to") or today.year)[:4])
        return [str(y) for y in range(y1, y0 - 1, -1)]          # 최신부터
    if kind == "month":
        y, m = (int(x) for x in str(it["from"]).split("-"))
        end = str(it.get("to") or f"{today:%Y-%m}")
        ey, em = (int(x) for x in end.split("-"))
        out = []
        while (y, m) <= (ey, em):
            out.append(f"{y}{m:02d}")
            y, m = (y + 1, 1) if m == 12 else (y, m + 1)
        return out[::-1]                                          # 최신부터
    raise ValueError(kind)


def fetch_mafra(src: dict, limit: int | None = None) -> Iterator[dict]:
    """농식품 빅데이터 포털: /openapi/{KEY}/json/{SID}/{start}/{end} (1000건 단위).

    iterate 가 있으면 그 파라미터(년월·연도 등)를 바꿔 가며 반복 호출한다.
    """
    sid = src["service_id"]
    combos = [{}]
    if src.get("iterate"):
        it = src["iterate"]
        combos = [{it["param"]: v} for v in iter_values(it)]
    seen = 0
    with httpx.Client(timeout=TIMEOUT) as c:
        for extra in combos:
            per, start = 1000, 1
            while True:
                end = start + per - 1
                url = f"{MAFRA_BASE}/{MAFRA_KEY}/json/{sid}/{start}/{end}"
                r = c.get(url, params={**(src.get("params") or {}), **extra})
                r.raise_for_status()
                data = r.json()
                save_raw(src["id"], data)
                body = data.get(sid, data)
                code = (body.get("result") or {}).get("code") if isinstance(body, dict) else None
                if code and code not in ("INFO-000", "INFO-0", "INFO-200"):
                    raise RuntimeError(f"농식품 포털 오류 {code}: {(body.get('result') or {}).get('message')}")
                recs = find_records(body)
                for rec in recs:
                    rec.update({f"_{k}": v for k, v in extra.items()})
                    yield rec
                    seen += 1
                    if limit and seen >= limit:
                        return
                total = total_count(body) or 0
                if not recs or end >= total:
                    break
                start = end + 1


def fetch_paged(src: dict, base_params: dict, limit: int | None = None) -> Iterator[dict]:
    """공공데이터포털 REST / odcloud 공통 페이지 순회."""
    pg = src.get("paging") or {"page": "pageNo", "size": "numOfRows", "per": 500}
    per = min(pg.get("per", 500), limit or 10**9)
    page = 1
    seen = 0
    with httpx.Client(timeout=TIMEOUT) as c:
        while True:
            params = {**base_params, **(src.get("params") or {}), pg["page"]: page, pg["size"]: per}
            r = c.get(src["endpoint"], params=params)
            if r.headers.get("content-type", "").startswith(("text/xml", "application/xml")) or r.text.lstrip().startswith("<"):
                raise RuntimeError(f"XML 응답(오류일 가능성): {r.text[:300]}")
            r.raise_for_status()
            data = r.json()
            save_raw(src["id"], data)
            head = json.dumps(data)[:400]
            if '"resultCode": "00"' not in head and '"resultCode"' in head and '"resultCode": "0"' not in head:
                m = re.search(r'"resultMsg"\s*:\s*"([^"]+)', head)
                if m and m.group(1) not in ("NORMAL SERVICE.", "NORMAL_SERVICE", "OK"):
                    raise RuntimeError(f"API 오류: {m.group(1)}")
            recs = find_records(data)
            yield from recs
            seen += len(recs)
            total = total_count(data)
            if not recs or (total is not None and seen >= total) or (limit and seen >= limit) or len(recs) < per:
                break
            page += 1


def fetch_file(src: dict, limit: int | None = None) -> Iterator[dict]:
    path = ROOT / src["path"]
    if path.suffix.lower() == ".csv":
        for enc in ("utf-8-sig", "cp949"):
            try:
                with open(path, encoding=enc, newline="") as f:
                    rows = list(csv.DictReader(f))
                break
            except UnicodeDecodeError:
                continue
    elif path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        ws = load_workbook(path, read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(next(it))]
        rows = [dict(zip(header, r)) for r in it if any(v is not None for v in r)]
    else:
        raise RuntimeError(f"지원하지 않는 파일 형식: {path.name}")
    yield from rows[:limit] if limit else rows


def fetch(src: dict, limit: int | None = None) -> Iterator[dict]:
    p = src["provider"]
    if p == "mafra":
        return fetch_mafra(src, limit)
    if p in ("datagokr", "odcloud"):
        return fetch_paged(src, {"serviceKey": DATA_GO_KR_KEY}, limit)
    if p == "file":
        return fetch_file(src, limit)
    if p == "kma_asos":
        return fetch_kma_asos(src, limit)
    raise ValueError(p)


def fetch_kma_asos(src: dict, limit: int | None = None) -> Iterator[dict]:
    end = date.today() - timedelta(days=1)
    start = end - timedelta(days=int(src.get("days", 60)))
    for stn, name in src["stations"].items():
        s = {**src, "params": {"dataType": "JSON", "dataCd": "ASOS", "dateCd": "DAY",
                               "startDt": f"{start:%Y%m%d}", "endDt": f"{end:%Y%m%d}", "stnIds": stn},
             "paging": {"page": "pageNo", "size": "numOfRows", "per": 999}}
        for rec in fetch_paged(s, {"serviceKey": DATA_GO_KR_KEY}, limit):
            rec["_stn_name"] = name
            yield rec


# ── 정규화 ──────────────────────────────────────────────────────────────
def normalize(src: dict, rec: dict) -> list[dict]:
    target, m, fixed = src["target"], src.get("map", {}), src.get("fixed", {})
    out: dict[str, Any] = {k: pick(rec, v) for k, v in m.items()}
    out.update({k: v for k, v in fixed.items() if out.get(k) in (None, "")})

    if src["provider"] == "kma_asos":
        return _asos_rows(rec)

    if src.get("addr_from"):                      # "전라남도 나주시 봉황면 덕림리" → 시도/시군구/읍면동
        parts = str(pick(rec, src["addr_from"]) or "").split()
        for key, i in (("sido", 0), ("sigungu", 1), ("emd", 2)):
            if not out.get(key) and len(parts) > i:
                out[key] = parts[i]
        if len(parts) > 3 and parts[2].endswith("구") and not out.get("emd_fixed"):
            out["sigungu"], out["emd"] = f"{parts[1]} {parts[2]}", parts[3]   # 시 + 구 2단 행정구역

    if src.get("metrics"):
        return _metric_rows(src, rec, out)

    if src.get("melt"):
        return _melt_scale(src, rec, out)

    lat = lon = None
    if "coords" in src:
        c = src["coords"]
        pt = geo.to_wgs84(pick(rec, c["x"]), pick(rec, c["y"]), str(c.get("crs", "4326")))
        if pt:
            lat, lon = pt

    if target == "event":
        row = {
            "kind": out.get("kind") or "outbreak",
            "disease": disease_code(out.get("disease")) or out.get("disease"),
            "species": out.get("species") if out.get("species") in ("poultry", "pig", "cattle", "wild_bird", "wild_boar")
            else species_code(str(out.get("species") or "") + " " + str(out.get("subspecies") or "")),
            "subspecies": out.get("subspecies") or (out.get("species") if out.get("species") not in ("poultry", "pig", "cattle") else None),
            "date": norm_date(out.get("date")),
            "sampled": norm_date(out.get("sampled")),
            "sido": out.get("sido"), "sigungu": out.get("sigungu"), "place": out.get("place"),
            "legal_code": str(out["legal_code"]) if out.get("legal_code") else None,
            "lat": lat, "lon": lon,
            "heads": int(num(out.get("heads")) or 0) or None,
            "note": out.get("note"),
        }
        if not row["date"]:
            return []
    elif target == "facility":
        row = {
            "kind": {"백신접종팀": "vaccine_team", "상하차반": "loading_crew"}.get(str(out.get("kind")), out.get("kind") or "farm"),
            "species": species_code(str(out.get("species") or "")) or out.get("species"),
            "subspecies": out.get("species") if out.get("species") not in ("poultry", "pig", "cattle") else None,
            "name": out.get("name"), "status": out.get("status"),
            "size": int(num(out.get("size")) or 0) or None,
            "sido": out.get("sido"), "sigungu": out.get("sigungu") or _sigungu_from_addr(out.get("address")),
            "address": out.get("address"), "lat": lat, "lon": lon,
        }
        if not row["sido"] and row["address"]:
            row["sido"] = str(row["address"]).split()[0]
    elif target == "contact":
        at_raw = f"{out.get('at') or ''} {out.get('at_time') or ''}"
        row = {
            "farm_key": out.get("farm_key"), "at": (norm_date(out.get("at")) or "") + _time_part(at_raw),
            "purpose": out.get("purpose"), "party": out.get("party"),
        }
        if not out.get("at"):
            return []
    elif target == "context_ts":
        metric = out.get("metric") or "value"
        if out.get("metric_suffix"):
            metric = f"{metric}:{out['metric_suffix']}"
        row = {"place": str(out.get("place") or "unknown"), "date": norm_date(out.get("date")) or "",
               "metric": metric, "value": num(out.get("value")) if out.get("value") is not None else 1.0,
               "lat": lat, "lon": lon}
        return [{**row, "source": src["id"]}] if row["date"] else []
    elif target == "region_stat":
        row = {"sido": out.get("sido") or "", "sigungu": out.get("sigungu") or "", "emd": out.get("emd") or "",
               "species": species_code(str(out.get("species") or "")) or out.get("species") or "",
               "metric": out.get("metric") or "value", "value": num(out.get("value")),
               "as_of": norm_date(out.get("as_of")) or ""}
        return [{**row, "source": src["id"]}]
    else:
        raise ValueError(target)

    keys = src.get("uid_from") or list(row)
    uid_parts = [row.get(k, rec.get(k)) for k in keys]
    row["uid"] = make_uid(src["id"], uid_parts)
    row["source"] = src["id"]
    row["raw"] = rec
    return [row]


def _metric_rows(src: dict, rec: dict, out: dict) -> list[dict]:
    """한 레코드의 여러 수치 컬럼을 metric 행 여러 개로 (region_stat / context_ts).

    metrics: [{metric: 이름, value: [컬럼], suffix: [컬럼](선택)}]
    """
    rows = []
    sp = species_code(str(out.get("species") or "")) or (out.get("species") or "")
    when = norm_date(out.get("as_of") or out.get("date")) or ""
    for mdef in src["metrics"]:
        val = num(pick(rec, mdef["value"]))
        if val is None:
            continue
        name = mdef["metric"]
        # suffix 후보 컬럼 값을 모두 이어 붙인다(예: 축종:질병) — 가금(닭·오리)처럼 묶이는 값이 덮어쓰이지 않도록
        sufs = [str(rec[c]).strip() for c in (mdef.get("suffix") or []) if rec.get(c) not in (None, "")]
        suf = ":".join(sufs) or out.get("metric_suffix")
        if suf:
            name = f"{name}:{suf}"
        if src["target"] == "region_stat":
            rows.append({"sido": out.get("sido") or "", "sigungu": out.get("sigungu") or "", "emd": out.get("emd") or "",
                         "species": sp, "metric": name, "value": val,
                         "as_of": (str(out.get("as_of")) if out.get("as_of") and len(re.sub(r"\D", "", str(out.get("as_of")))) <= 4 else when),
                         "source": src["id"]})
        else:
            if not when:
                continue
            rows.append({"place": str(out.get("place") or out.get("sido") or "전국"), "date": when, "metric": name,
                         "value": val, "lat": None, "lon": None, "source": src["id"]})
    return rows


def _time_part(v: Any) -> str:
    m = re.search(r"(\d{1,2}):(\d{2})", str(v or ""))
    return f" {int(m.group(1)):02d}:{m.group(2)}" if m else ""


def _sigungu_from_addr(addr: Any) -> str | None:
    if not addr:
        return None
    parts = str(addr).split()
    return parts[1] if len(parts) > 1 else None


def _asos_rows(rec: dict) -> list[dict]:
    d = norm_date(rec.get("tm"))
    if not d:
        return []
    place = f"{rec.get('stnId')}:{rec.get('_stn_name') or rec.get('stnNm')}"
    pairs = {"temp_avg": "avgTa", "temp_min": "minTa", "humidity_avg": "avgRhm",
             "wind_dir_max": "maxWsWd", "wind_speed_avg": "avgWs", "rain": "sumRn"}
    return [{"place": place, "date": d, "metric": k, "value": num(rec.get(v)), "lat": None, "lon": None,
             "source": "kma_asos_daily"} for k, v in pairs.items() if num(rec.get(v)) is not None]


def _melt_scale(src: dict, rec: dict, out: dict) -> list[dict]:
    """'50000수 이상 호수', '... 두수' 같은 넓은 표를 metric 행으로 펼친다."""
    rows = []
    skip = {c.lower() for cands in src.get("map", {}).values() for c in cands} | {c.lower() for c in src.get("melt_skip", [])}
    for k, v in rec.items():
        if k.lower() in skip:
            continue
        val = num(v)
        if val is None:
            continue
        metric = "farms" if "호" in k else "heads" if ("두" in k or "수수" in k) else k
        rows.append({"sido": out.get("sido") or "", "sigungu": out.get("sigungu") or "",
                     "species": src.get("fixed", {}).get("species") or "",
                     "metric": f"{metric}:{k}", "value": val,
                     "as_of": norm_date(out.get("as_of")) or "", "source": src["id"]})
    return rows


# ── 실행 ────────────────────────────────────────────────────────────────
def run(sid: str, limit: int | None = None) -> int:
    """수집 → 원본 레코드(raw_record) 저장 → 정규화 테이블 upsert → collect_log 기록."""
    return run_source(get_source(sid), limit)


def run_source(src: dict, limit: int | None = None) -> int:
    if src["provider"] == "birds":
        from . import birds
        return birds.load()["rows"]
    if src["provider"] == "kahis_visit":          # 자체 수집기(날짜·시도 단위, 기록도 스스로 남김)
        from . import kahis_visit
        return kahis_visit.run()["visits"]       # 매일분(어제·오늘). 과거분은 kahis_visit 백그라운드 스레드
    sid = src["id"]
    started = db.now()
    batch_id = f"{started:%Y%m%d%H%M%S}"
    try:
        if not is_configured(src):
            raise RuntimeError("설정 안 됨: endpoint/service_id/키/파일 경로를 확인하세요")
        raws: list[dict] = []
        rows: list[dict] = []
        for rec in fetch(src, limit):
            raws.append(rec)
            rows.extend(normalize(src, rec))
        with db.session() as con:
            db.save_raw(con, sid, batch_id, raws)
            if src["target"] == "region_stat" and not src.get("_recent"):
                # 전체 수집이면 그 원천의 집계를 통째로 바꾼다(사라진 지역·항목이 남지 않게)
                con.execute(f"DELETE FROM region_stat WHERE source = ?", [sid])
            n = db.upsert(con, src["target"], rows)
            con.con.execute(db.collect_log.insert().values(source=sid, started_at=started, finished_at=db.now(),
                                                           ok=True, rows=n, message=f"원본 {len(raws)}건"))
        return n
    except Exception as e:  # 수집 상태 화면에 남긴다
        with db.session() as con:
            con.con.execute(db.collect_log.insert().values(source=sid, started_at=started, finished_at=db.now(),
                                                           ok=False, rows=0, message=str(e)[:500]))
        raise


def probe(sid: str, n: int = 3) -> dict:
    """정규화 없이 원본 레코드 몇 건과 필드 목록만 본다."""
    src = get_source(sid)
    recs = []
    for rec in fetch(src, n):
        recs.append(rec)
        if len(recs) >= n:
            break
    fields = sorted({k for r in recs for k in r})
    return {"source": sid, "fields": fields, "sample": recs,
            "normalized": [normalize(src, r) for r in recs]}
